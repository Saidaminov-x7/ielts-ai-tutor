import json
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook  # type: ignore[import-untyped]
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.database.models import CourseDay, Vocabulary

TARGET_NAME = "IELTS_7_5_250_istoriy_26_nedel_182_dnya_PRO.xlsx"


@dataclass
class ImportReport:
    sheets: int = 0
    days_added: int = 0
    days_updated: int = 0
    words_added: int = 0
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    def save(self, path: Path) -> None:
        path.write_text(json.dumps(self.__dict__, ensure_ascii=False, indent=2), encoding="utf-8")


def find_source(path: Path | None = None) -> Path | None:
    if path:
        return path if path.is_file() else None
    candidates = [Path.cwd() / TARGET_NAME, Path.cwd() / "data" / TARGET_NAME]
    return next((item for item in candidates if item.is_file()), None)


def normalize(value: object) -> str:
    return str(value).strip() if value is not None else ""


def cell_value(row: tuple, indices: dict[str, int], keys: tuple[str, ...]) -> str:
    position = next((indices[key] for key in keys if key in indices), -1)
    return normalize(row[position]) if 0 <= position < len(row) else ""


def import_workbook(session: Session, path: Path) -> ImportReport:
    report = ImportReport()
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        report.errors.append(f"Не удалось открыть XLSX: {type(exc).__name__}")
        return report
    report.sheets = len(workbook.sheetnames)
    day_count = word_count = 0
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            report.warnings.append(f"Лист '{sheet.title}': пустой")
            continue
        names = [normalize(value).casefold() for value in headers]
        if any(name in names for name in ("day", "день", "day_number")):
            idx = {name: i for i, name in enumerate(names) if name}
            day_idx = next(idx[n] for n in ("day", "день", "day_number") if n in idx)
            for row_num, row in enumerate(rows, 2):
                if day_idx >= len(row) or not isinstance(row[day_idx], (int, float)):
                    continue
                day = int(row[day_idx])
                if day < 1 or day > 182:
                    report.errors.append(f"{sheet.title}:{row_num}: день вне диапазона 1–182")
                    continue

                topic = cell_value(row, idx, ("topic", "тема", "theme", "название"))
                if not topic:
                    report.errors.append(f"{sheet.title}:{row_num}: отсутствует тема дня {day}")
                    continue
                record = session.scalar(select(CourseDay).where(CourseDay.day_number == day))
                values = dict(
                    week_number=(day - 1) // 7 + 1,
                    topic=topic,
                    grammar=cell_value(row, idx, ("grammar", "грамматика")),
                    reading=cell_value(row, idx, ("reading", "чтение")),
                    listening=cell_value(row, idx, ("listening", "аудирование")),
                    writing=cell_value(row, idx, ("writing", "письмо")),
                    objectives=cell_value(row, idx, ("objectives", "цели", "задания")),
                )
                if record:
                    for key, value in values.items():
                        setattr(record, key, value)
                    report.days_updated += 1
                else:
                    session.add(CourseDay(day_number=day, **values))
                    report.days_added += 1
                day_count += 1
        if any(name in names for name in ("english", "word", "слово", "английское слово")):
            idx = {name: i for i, name in enumerate(names) if name}
            en_i = next(
                idx[n] for n in ("english", "word", "слово", "английское слово") if n in idx
            )
            ru_i = next((idx[n] for n in ("translation", "перевод", "русский") if n in idx), -1)
            if ru_i < 0:
                report.errors.append(
                    f"Лист '{sheet.title}': найден список слов без столбца перевода"
                )
                continue
            for row_num, row in enumerate(rows, 2):
                en = normalize(row[en_i]) if en_i < len(row) else ""
                ru = normalize(row[ru_i]) if ru_i < len(row) else ""
                if not en and not ru:
                    continue
                if not en or not ru:
                    report.errors.append(f"{sheet.title}:{row_num}: неполная запись словаря")
                    continue
                existing = session.scalar(
                    select(Vocabulary).where(Vocabulary.english == en, Vocabulary.translation == ru)
                )
                if existing:
                    continue
                course_day_raw = cell_value(row, idx, ("day", "day_number", "день"))
                course_day = (
                    int(course_day_raw)
                    if course_day_raw.isdigit() and 1 <= int(course_day_raw) <= 182
                    else None
                )

                session.add(
                    Vocabulary(
                        english=en,
                        translation=ru,
                        transcription=cell_value(row, idx, ("transcription", "транскрипция")),
                        example=cell_value(row, idx, ("example", "пример")),
                        topic=cell_value(row, idx, ("topic", "тема")),
                        course_day=course_day,
                    )
                )
                report.words_added += 1
                word_count += 1
    if day_count == 0:
        report.warnings.append("Не найдено распознанных строк с учебными днями")
    if word_count == 0:
        report.warnings.append("Не найдено распознанных строк словаря")
    if report.days_added + report.days_updated not in (0, 182):
        report.warnings.append(
            f"Импортировано данных о {report.days_added + report.days_updated} днях; ожидается 182"
        )
    session.commit()
    workbook.close()
    return report
